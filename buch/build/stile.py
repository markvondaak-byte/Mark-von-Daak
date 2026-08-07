"""Satzspiegel, Absatzformate und Word-Feldfunktionen für beide Bände.

Wird von build_docx.py (Band 1) und workbook/build/build_workbook.py (Band 2)
importiert, damit Typografie und Farben nur an einer Stelle definiert sind.
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# --- Schriften ---------------------------------------------------------------
# Liberation ist die einzige im Container verfügbare Familie mit allen vier
# Schnitten als TrueType (Regular/Kursiv/Fett/Fett-Kursiv). TrueType ist
# Voraussetzung dafür, dass LibreOffice die Schrift ins PDF einbettet — KDP
# lehnt PDFs mit nicht eingebetteten Schriften ab.
SERIF = "Liberation Serif"
SANS = "Liberation Sans"

# --- Farben ------------------------------------------------------------------
# Frische, helle Palette. Bewusst NICHT das FitLine-Crimson #C8102E der
# Website, damit das Buch nicht wie eine Publikation von PM-International wirkt.
FARBEN = {
    "text": RGBColor(0x1A, 0x1A, 0x1A),
    "gedaempft": RGBColor(0x5A, 0x5A, 0x5A),
    "blatt": RGBColor(0x2F, 0x7A, 0x3E),  # Leitfarbe, frisches Blattgrün
    "blatt_hell": RGBColor(0x6F, 0xA8, 0x5E),
    "linie": RGBColor(0xC8, 0xC8, 0xC8),
    "kasten": "EEF4EA",  # Schattierung als Hex-String (Word erwartet das so)
    "warnung": "FBF0E4",
}

# Tagesfarben des Konzepts. Für Band 2 wichtig: der Druck erfolgt in
# Schwarz-Weiß, deshalb trägt jede Farbe zusätzlich eine Beschriftung.
TAGESFARBEN = {
    "weiss": {"label": "WEISS", "fuellung": "FFFFFF", "rahmen": "9A9A9A"},
    "gruen": {"label": "GRÜN", "fuellung": "DCEBD2", "rahmen": "6FA85E"},
    "rot": {"label": "ROT", "fuellung": "F5DCDC", "rahmen": "B5564F"},
    "vorbereitung": {"label": "VORBEREITUNG", "fuellung": "ECECEC", "rahmen": "9A9A9A"},
}


# --- Dokument und Satzspiegel ------------------------------------------------
def dokument_anlegen(seitenformat):
    """Erzeugt ein Dokument mit KDP-Satzspiegel und allen Absatzformaten.

    `seitenformat` ist das Dict aus buch.yaml / workbook.yaml.
    """
    doc = Document()
    _leeren(doc)
    _spiegelraender_aktivieren(doc)
    _gerade_ungerade_kopfzeilen(doc)
    seite_einrichten(doc.sections[0], seitenformat)
    _stile_definieren(doc)
    return doc


def _leeren(doc):
    """Entfernt den leeren Standardabsatz, den python-docx mitliefert."""
    for absatz in list(doc.paragraphs):
        absatz._element.getparent().remove(absatz._element)


def seite_einrichten(section, sf):
    section.page_width = Mm(sf["breite_mm"])
    section.page_height = Mm(sf["hoehe_mm"])
    section.top_margin = Mm(sf["rand_oben_mm"])
    section.bottom_margin = Mm(sf["rand_unten_mm"])
    # Bei gespiegelten Rändern ist links = innen und rechts = außen.
    section.left_margin = Mm(sf["rand_innen_mm"])
    section.right_margin = Mm(sf["rand_aussen_mm"])
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)
    return section


def _einstellung_setzen(doc, name, vor_elementen):
    """Fügt ein Schalter-Element in word/settings.xml ein.

    CT_Settings ist im OOXML-Schema eine feste Sequenz. Ein einfaches append()
    landet hinter <w:compat>/<w:rsids> und macht die Datei ungültig —
    LibreOffice bricht den Import dann mit "source file could not be loaded" ab.
    Deshalb wird vor dem ersten passenden Anker eingefügt.
    """
    einstellungen = doc.settings.element
    if einstellungen.find(qn(name)) is not None:
        return
    element = OxmlElement(name)
    for anker_name in vor_elementen:
        anker = einstellungen.find(qn(anker_name))
        if anker is not None:
            anker.addprevious(element)
            return
    einstellungen.append(element)


def _spiegelraender_aktivieren(doc):
    """<w:mirrorMargins/> — innen/außen statt links/rechts. Pflicht für Druck."""
    _einstellung_setzen(doc, "w:mirrorMargins",
                        ("w:proofState", "w:defaultTabStop", "w:compat"))


def _gerade_ungerade_kopfzeilen(doc):
    _einstellung_setzen(doc, "w:evenAndOddHeaders",
                        ("w:characterSpacingControl", "w:compat"))


# --- Absatzformate -----------------------------------------------------------
def _stil(doc, name, *, schrift, groesse, fett=False, kursiv=False,
          farbe=None, vor=0, nach=0, zeilen=1.15, ausrichtung=None,
          zusammenhalten=False, einzug_links=None):
    stil = doc.styles.add_style(name, 1)  # 1 = WD_STYLE_TYPE.PARAGRAPH
    stil.quick_style = True
    schriftbild = stil.font
    schriftbild.name = schrift
    schriftbild.size = Pt(groesse)
    schriftbild.bold = fett
    schriftbild.italic = kursiv
    schriftbild.color.rgb = farbe or FARBEN["text"]
    # Ostasiatische und Komplexschrift-Varianten mitsetzen, sonst fällt Word
    # bei Umlauten gelegentlich auf eine Ersatzschrift zurück.
    rpr = stil.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), schrift)

    absatz = stil.paragraph_format
    absatz.space_before = Pt(vor)
    absatz.space_after = Pt(nach)
    absatz.line_spacing = zeilen
    absatz.keep_with_next = zusammenhalten
    absatz.widow_control = True
    if ausrichtung is not None:
        absatz.alignment = ausrichtung
    if einzug_links is not None:
        absatz.left_indent = Mm(einzug_links)
    return stil


def _stile_definieren(doc):
    mitte = WD_ALIGN_PARAGRAPH.CENTER

    # Titelei
    _stil(doc, "BuchTitel", schrift=SANS, groesse=26, fett=True,
          farbe=FARBEN["blatt"], vor=0, nach=8, zeilen=1.0, ausrichtung=mitte)
    _stil(doc, "BuchUntertitel", schrift=SANS, groesse=14,
          farbe=FARBEN["gedaempft"], vor=0, nach=24, zeilen=1.2, ausrichtung=mitte)
    _stil(doc, "BuchAutor", schrift=SANS, groesse=13, fett=True,
          vor=0, nach=0, ausrichtung=mitte)

    # Gliederung
    _stil(doc, "Teilnummer", schrift=SANS, groesse=11, fett=True,
          farbe=FARBEN["blatt_hell"], vor=0, nach=4, ausrichtung=mitte)
    _stil(doc, "Teiltitel", schrift=SANS, groesse=22, fett=True,
          farbe=FARBEN["blatt"], vor=0, nach=0, zeilen=1.1, ausrichtung=mitte)
    _stil(doc, "KapitelNummer", schrift=SANS, groesse=10, fett=True,
          farbe=FARBEN["blatt_hell"], vor=0, nach=3)
    _stil(doc, "Kapitel", schrift=SANS, groesse=19, fett=True,
          farbe=FARBEN["blatt"], vor=0, nach=10, zeilen=1.1, zusammenhalten=True)
    _stil(doc, "Abschnitt", schrift=SANS, groesse=13, fett=True,
          vor=13, nach=5, zeilen=1.15, zusammenhalten=True)
    _stil(doc, "Unterabschnitt", schrift=SANS, groesse=10.5, fett=True,
          farbe=FARBEN["gedaempft"], vor=8, nach=3, zusammenhalten=True)

    # Fließtext
    _stil(doc, "Fliesstext", schrift=SERIF, groesse=11, vor=0, nach=6,
          zeilen=1.15)
    _stil(doc, "FliesstextEng", schrift=SERIF, groesse=11, vor=0, nach=2,
          zeilen=1.15)
    _stil(doc, "Einzug", schrift=SERIF, groesse=11, vor=0, nach=5,
          zeilen=1.10, einzug_links=8)
    _stil(doc, "Klein", schrift=SERIF, groesse=9, farbe=FARBEN["gedaempft"],
          vor=0, nach=5)
    _stil(doc, "KleinMitte", schrift=SERIF, groesse=9, farbe=FARBEN["gedaempft"],
          vor=0, nach=5, ausrichtung=mitte)
    _stil(doc, "Zitat", schrift=SERIF, groesse=11, kursiv=True,
          farbe=FARBEN["gedaempft"], vor=6, nach=10, einzug_links=6)

    # Listen
    _stil(doc, "Punkt", schrift=SERIF, groesse=11, vor=0, nach=2,
          zeilen=1.10, einzug_links=6)
    _stil(doc, "Nummer", schrift=SERIF, groesse=11, vor=0, nach=2,
          zeilen=1.10, einzug_links=6)

    # Kästen und Tabellen
    _stil(doc, "KastenTitel", schrift=SANS, groesse=10, fett=True,
          farbe=FARBEN["blatt"], vor=0, nach=3, zusammenhalten=True)
    _stil(doc, "KastenText", schrift=SERIF, groesse=9.5, vor=0, nach=3,
          zeilen=1.15)
    _stil(doc, "TabellenKopf", schrift=SANS, groesse=9, fett=True,
          vor=2, nach=2, zeilen=1.05)
    _stil(doc, "TabellenZelle", schrift=SERIF, groesse=9, vor=2, nach=2,
          zeilen=1.05)

    # Inhaltsverzeichnis
    # Enger als der Fließtext: 32 Einträge sollen auf eine Seite passen.
    # Sonst steht der letzte Eintrag allein auf einer zweiten Seite und
    # sieht nach einem Satzfehler aus.
    _stil(doc, "InhaltTeil", schrift=SANS, groesse=10, fett=True,
          farbe=FARBEN["blatt"], vor=8, nach=2)
    _stil(doc, "InhaltKapitel", schrift=SERIF, groesse=10, vor=0, nach=1.5)

    # Kopf- und Fußzeile
    _stil(doc, "Kopfzeile", schrift=SANS, groesse=8, farbe=FARBEN["gedaempft"],
          vor=0, nach=0)
    _stil(doc, "Fusszeile", schrift=SANS, groesse=9, farbe=FARBEN["gedaempft"],
          vor=0, nach=0)

    # Standardschrift des Dokuments, damit auch nicht formatierte Läufe passen
    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(11)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), SERIF)


# --- Word-Felder -------------------------------------------------------------
def feld_einfuegen(absatz, anweisung):
    """Fügt ein Word-Feld ein (z. B. 'PAGE' oder 'TOC \\o "1-2" \\h \\z \\u').

    Felder sind nötig, weil Seitenzahlen und Inhaltsverzeichnis erst beim
    Rendern berechnet werden — LibreOffice löst sie beim PDF-Export auf.
    """
    lauf = absatz.add_run()

    anfang = OxmlElement("w:fldChar")
    anfang.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = anweisung
    trennung = OxmlElement("w:fldChar")
    trennung.set(qn("w:fldCharType"), "separate")
    platzhalter = OxmlElement("w:t")
    platzhalter.text = "…"
    ende = OxmlElement("w:fldChar")
    ende.set(qn("w:fldCharType"), "end")

    for teil in (anfang, text, trennung, platzhalter, ende):
        lauf._r.append(teil)
    return absatz


def tabulator_rechts(absatz, position_mm, fuellzeichen="dot"):
    """Rechtsbündiger Tabstopp mit Füllpunkten — für Inhaltsverzeichniszeilen."""
    ppr = absatz._p.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        ppr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), fuellzeichen)
    tab.set(qn("w:pos"), str(int(Mm(position_mm).twips)))
    tabs.append(tab)
    return absatz


# --- Kopf- und Fußzeilen -----------------------------------------------------
def kopf_und_fusszeile(doc, section, links_text, rechts_text,
                       *, mit_seitenzahl=True):
    """Setzt Kopfzeile (Buchtitel / Kapitel) und Fußzeile (Seitenzahl).

    Bei gespiegeltem Satz steht die Seitenzahl außen: auf ungeraden (rechten)
    Seiten rechts, auf geraden (linken) Seiten links.

    `doc` wird gebraucht, weil Kopf-/Fußzeilen eigene Parts sind und von dort
    kein Weg zurück zu den Dokumentstilen führt.
    """
    section.different_first_page_header_footer = True

    def _fuellen(kopfzeile, fusszeile, text, ausrichtung_zahl):
        kopfzeile.is_linked_to_previous = False
        fusszeile.is_linked_to_previous = False
        _ersten_absatz_leeren(kopfzeile)
        _ersten_absatz_leeren(fusszeile)

        k = kopfzeile.paragraphs[0]
        k.style = doc.styles["Kopfzeile"]
        k.alignment = ausrichtung_zahl
        k.add_run(text)

        if mit_seitenzahl:
            f = fusszeile.paragraphs[0]
            f.style = doc.styles["Fusszeile"]
            f.alignment = ausrichtung_zahl
            feld_einfuegen(f, "PAGE")

    # Ungerade = rechte Seite → außen ist rechts
    _fuellen(section.header, section.footer, rechts_text,
             WD_ALIGN_PARAGRAPH.RIGHT)
    # Gerade = linke Seite → außen ist links
    _fuellen(section.even_page_header, section.even_page_footer, links_text,
             WD_ALIGN_PARAGRAPH.LEFT)

    # Erste Kapitelseite bleibt ohne Kopfzeile, trägt aber die Seitenzahl.
    section.first_page_header.is_linked_to_previous = False
    _ersten_absatz_leeren(section.first_page_header)
    section.first_page_footer.is_linked_to_previous = False
    _ersten_absatz_leeren(section.first_page_footer)
    if mit_seitenzahl:
        f = section.first_page_footer.paragraphs[0]
        f.style = doc.styles["Fusszeile"]
        f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        feld_einfuegen(f, "PAGE")


def _ersten_absatz_leeren(bereich):
    if not bereich.paragraphs:
        bereich.add_paragraph()
    absatz = bereich.paragraphs[0]
    for lauf in list(absatz.runs):
        lauf._r.getparent().remove(lauf._r)
    return absatz


def seitenzahlen_format(section, format_="decimal", neustart_bei=None):
    """Setzt <w:pgNumType>, z. B. römische Ziffern für die Titelei."""
    sect_pr = section._sectPr
    vorhanden = sect_pr.find(qn("w:pgNumType"))
    if vorhanden is not None:
        sect_pr.remove(vorhanden)
    element = OxmlElement("w:pgNumType")
    element.set(qn("w:fmt"), format_)
    if neustart_bei is not None:
        element.set(qn("w:start"), str(neustart_bei))
    sect_pr.append(element)


# --- Bausteine ---------------------------------------------------------------
def seitenumbruch(doc):
    absatz = doc.add_paragraph()
    absatz.add_run().add_break(WD_BREAK.PAGE)
    return absatz


def neuer_abschnitt(doc, seitenformat, *, ungerade_start=False):
    """Neuer Word-Abschnitt — nötig, wenn sich Kopfzeile oder
    Seitenzahlformat ändern."""
    art = WD_SECTION.ODD_PAGE if ungerade_start else WD_SECTION.NEW_PAGE
    section = doc.add_section(art)
    seite_einrichten(section, seitenformat)
    return section


def schattierung(zelle_oder_absatz, hexfarbe):
    """Hintergrundfarbe für Tabellenzelle oder Absatz."""
    element = getattr(zelle_oder_absatz, "_tc", None)
    if element is not None:
        ziel = element.get_or_add_tcPr()
    else:
        ziel = zelle_oder_absatz._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfarbe)
    ziel.append(shd)


def rahmen(absatz, hexfarbe, *, staerke=6, seiten=("top", "bottom", "left", "right")):
    """Rahmenlinien um einen Absatz — für Hinweiskästen."""
    ppr = absatz._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for seite in seiten:
        linie = OxmlElement(f"w:{seite}")
        linie.set(qn("w:val"), "single")
        linie.set(qn("w:sz"), str(staerke))
        linie.set(qn("w:space"), "6")
        linie.set(qn("w:color"), hexfarbe)
        bdr.append(linie)
    ppr.append(bdr)


def tabellenraender(tabelle, hexfarbe="C8C8C8", staerke=4):
    """Dünne Linien innen und außen — python-docx bietet dafür nichts an."""
    tbl_pr = tabelle._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for kante in ("top", "left", "bottom", "right", "insideH", "insideV"):
        linie = OxmlElement(f"w:{kante}")
        linie.set(qn("w:val"), "single")
        linie.set(qn("w:sz"), str(staerke))
        linie.set(qn("w:space"), "0")
        linie.set(qn("w:color"), hexfarbe)
        borders.append(linie)
    tbl_pr.append(borders)


def spaltenbreiten(tabelle, breiten_mm):
    """Erzwingt feste Spaltenbreiten.

    `cell.width` allein wirkt nicht — Word verteilt die Spalten sonst nach
    Inhalt. Nötig sind zusätzlich <w:tblLayout w:type="fixed"/> und ein
    passender <w:tblGrid>.
    """
    tbl = tabelle._tbl
    tbl_pr = tbl.tblPr

    vorhanden = tbl_pr.find(qn("w:tblLayout"))
    if vorhanden is not None:
        tbl_pr.remove(vorhanden)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for breite in breiten_mm:
        spalte = OxmlElement("w:gridCol")
        spalte.set(qn("w:w"), str(int(Mm(breite).twips)))
        grid.append(spalte)
    tbl.insert(tbl.index(tbl_pr) + 1, grid)

    for zeile in tabelle.rows:
        for zelle, breite in zip(zeile.cells, breiten_mm):
            zelle.width = Mm(breite)
    return tabelle


def zellenrand(tabelle, oben=1.0, unten=1.0, links=1.5, rechts=1.5):
    """Innenabstand aller Zellen in Millimetern."""
    tbl_pr = tabelle._tbl.tblPr
    rand = OxmlElement("w:tblCellMar")
    for name, wert in (("top", oben), ("bottom", unten),
                       ("left", links), ("right", rechts)):
        element = OxmlElement(f"w:{name}")
        element.set(qn("w:w"), str(int(Mm(wert).twips)))
        element.set(qn("w:type"), "dxa")
        rand.append(element)
    tbl_pr.append(rand)
    return tabelle


def zeile_zusammenhalten(zeile, ganz=True):
    """Verhindert, dass eine Tabellenzeile über den Seitenrand bricht."""
    tr_pr = zeile._tr.get_or_add_trPr()
    if ganz:
        tr_pr.append(OxmlElement("w:cantSplit"))


def kopfzeile_wiederholen(zeile):
    tr_pr = zeile._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
